from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"E:\Obisian Project\Lantern-Project")
SOURCE = ROOT / "implementation" / "competitive-research-2026-08-28" / "report-source.md"
OUTPUT = ROOT / "dist" / "Nimvara-Competitive-Complaints-and-Backlog-2026-08-28.docx"

NAVY = "17324D"
BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
WHITE = "FFFFFF"
GOLD = "A36B00"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[i]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline_markup(paragraph, text, size=10.5, color="222222"):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`|https?://\S+)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=size, color=color, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, size=size - 0.5, color=NAVY, name="Consolas")
        else:
            run = paragraph.add_run(part)
            set_font(run, size=size, color=color)


def add_body(doc, text, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    add_inline_markup(p, text)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    add_inline_markup(p, text)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    add_inline_markup(p, text)


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    rows = [r + [""] * (cols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    if cols == 2:
        widths = [2600, 6760]
    elif cols == 3:
        widths = [2100, 3500, 3760]
    elif cols == 4:
        widths = [1500, 3450, 1750, 2660]
    else:
        base = 9360 // cols
        widths = [base] * cols
        widths[-1] += 9360 - sum(widths)
    set_table_widths(table, widths)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if r_idx == 0:
                set_cell_shading(cell, NAVY)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(value)
            set_font(run, size=8.2 if cols >= 4 else 9, color=WHITE if r_idx == 0 else "222222", bold=r_idx == 0)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 9),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 11.5, NAVY, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_cover(doc):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("COMPETITIVE RESEARCH")
    set_font(r, size=10, color=GOLD, bold=True)
    p.paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Nimvara")
    set_font(r, size=30, color=NAVY, bold=True)
    p.paragraph_format.space_after = Pt(5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Competitive complaints research\nand product backlog")
    set_font(r, size=18, color=BLUE)
    p.paragraph_format.space_after = Pt(28)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Trustworthiness over feature breadth")
    set_font(r, size=11, color=MID_GRAY, italic=True)
    p.paragraph_format.space_after = Pt(70)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Evidence-backed desk research | 28 August 2026")
    set_font(r, size=10, color=MID_GRAY)
    doc.add_page_break()


def add_front_matter(doc):
    p = doc.add_paragraph()
    r = p.add_run("Research note")
    set_font(r, size=13, color=NAVY, bold=True)
    p.paragraph_format.space_after = Pt(5)
    add_body(doc, "This report synthesizes official documentation, current issue trackers, product forums, independent reporting, and clearly labeled community signals. It is not user research, a representative survey, or a measured competitor benchmark.")
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    values = [
        ("Canonical source", "implementation/competitive-research-2026-08-28/report-source.md"),
        ("Claim ledger", "implementation/competitive-research-2026-08-28/claim-source-ledger.md"),
        ("Decision owner", "Barry/Kogu"),
        ("Recommended next gate", "NIM-001 Workspace Safety Center + NIM-002 sync harness"),
    ]
    for i, (label, value) in enumerate(values):
        set_cell_shading(table.cell(i, 0), LIGHT_BLUE)
        for c, text in enumerate((label, value)):
            p = table.cell(i, c).paragraphs[0]
            r = p.add_run(text)
            set_font(r, size=9.3, color=NAVY if c == 0 else "222222", bold=c == 0)
            set_cell_margins(table.cell(i, c))
    set_table_widths(table, [2450, 6910])
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_page_break()


def parse_markdown(doc, text):
    lines = text.splitlines()
    i = 0
    table_rows = []
    skipped_title = False
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
                table_rows.append(cells)
            i += 1
            if i >= len(lines) or not lines[i].startswith("|"):
                add_table(doc, table_rows)
                table_rows = []
            continue
        if not line:
            i += 1
            continue
        if line.startswith("# ") and not skipped_title:
            skipped_title = True
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif re.match(r"^\d+\. ", line):
            add_number(doc, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            add_bullet(doc, line[2:])
        elif line.startswith("**") and line.endswith("  "):
            add_body(doc, line[:-2])
        else:
            add_body(doc, line)
        i += 1


def add_header_footer(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)
        hp = section.header.paragraphs[0]
        hp.text = "NIMVARA  |  COMPETITIVE RESEARCH"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_font(hp.runs[0], size=8, color=MID_GRAY, bold=True)
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("28 August 2026  |  Evidence-backed desk research")
        set_font(r, size=8, color=MID_GRAY)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    add_header_footer(doc)
    add_cover(doc)
    add_front_matter(doc)
    parse_markdown(doc, SOURCE.read_text(encoding="utf-8"))
    props = doc.core_properties
    props.title = "Nimvara Competitive Complaints Research and Product Backlog"
    props.subject = "Evidence-backed product strategy and prioritized engineering backlog"
    props.author = "Nimvara Project"
    props.keywords = "Nimvara, competitive research, knowledge management, product backlog"
    props.comments = "Derived from canonical Markdown source; no invented interviews or benchmarks."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
